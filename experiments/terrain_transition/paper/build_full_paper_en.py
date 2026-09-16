#!/usr/bin/env python3
"""English build of the full paper docx, in ONE pass with python-docx directly
(same rationale as build_full_paper.py: tables/figures must land inline in
their subsection, and docx_builder.py's merge drops image relationships).

This is the English translation of build_full_paper.py. It reads translated
copy from sections_en/ and tables_en/, keeps the same style names and layout,
and writes out/full_paper_en.docx. The corresponding-author email line is
intentionally omitted here.
"""
import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

ROOT = Path("/home/haktan/Projects/LunarRocket/experiments/terrain_transition/paper")
SEC = ROOT / "sections_en"
TBL = ROOT / "tables_en"
RES = ROOT.parent / "results"
RES_TC = RES / "training_comparison"
TPL = ROOT / "generic_template.docx"
OUT = ROOT / "out" / "full_paper_en.docx"


def split_body_text(raw):
    for block in raw.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("- "):
            yield "bullet", block[2:].strip()
        else:
            yield "para", " ".join(line.strip() for line in block.splitlines())


def add_title_block(doc, title, authors, affiliations, abstract_path, keywords):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(authors)
    r.font.size = Pt(12)

    for aff in affiliations:
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(aff)
        r.italic = True
        r.font.size = Pt(10)

    doc.add_paragraph(style="Normal")
    doc.add_paragraph(style="Normal")

    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract -")
    r.bold = True
    r.italic = True
    r.font.size = Pt(12)
    p.add_run(" ")
    abstract_text = Path(abstract_path).read_text(encoding="utf-8").strip()
    r = p.add_run(abstract_text)
    r.font.size = Pt(12)
    n_words = len(abstract_text.split())
    print(f"abstract word count: {n_words}")

    doc.add_paragraph(style="Normal")

    p = doc.add_paragraph(style="Normal")
    r = p.add_run("Keywords -")
    r.italic = True
    r.font.size = Pt(10)
    p.add_run(" ")
    r = p.add_run(", ".join(keywords))
    r.italic = True
    r.font.size = Pt(10)


def add_section(doc, heading, heading_style, body_style, path):
    if heading:
        doc.add_paragraph(heading, style=heading_style)
    raw = Path(path).read_text(encoding="utf-8")
    for kind, text in split_body_text(raw):
        prefix = "•  " if kind == "bullet" else ""
        doc.add_paragraph(prefix + text, style=body_style)


def add_table(doc, csv_path, caption, caption_style="Caption", table_style="Table Grid"):
    doc.add_paragraph(caption, style=caption_style)
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = table_style
    except KeyError:
        pass
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val


def add_figure(doc, image_path, caption, caption_style="Caption", width_inches=6.2):
    doc.add_picture(str(image_path), width=Inches(width_inches) if width_inches else None)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(caption, style=caption_style)


def add_references(doc, path):
    from docx.shared import Cm
    doc.add_paragraph("References", style="Heading 1")
    raw = Path(path).read_text(encoding="utf-8")
    for block in raw.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        text = " ".join(line.strip() for line in block.splitlines())
        p = doc.add_paragraph(text, style="Normal")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)


def main():
    doc = Document(str(TPL))
    body = doc.element.body
    from docx.oxml.ns import qn
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    title = ("Terrain Fine-Detail Transition in Lunar Landing Simulations: "
             "A Comparison of Smoothstep, Gaussian, and Hybrid Blending "
             "Methods in Isaac Lab")

    add_title_block(
        doc,
        title=title,
        authors="Haktan Yusuf Guloglu",
        affiliations=["Sakarya University, National Technology Workshop, Artificial Intelligence Center"],
        abstract_path=ROOT / "abstract_en.txt",
        keywords=["Lunar landing simulation", "Reinforcement learning", "Isaac Lab",
                  "Procedural terrain generation", "Multi-resolution blending",
                  "Continuity analysis", "Soft Actor-Critic"],
    )

    add_section(doc, "Abstract (Turkish)", "Heading 1", "Normal", SEC / "01_abstract_tr.txt")
    add_section(doc, "1. Introduction", "Heading 1", "Normal", SEC / "02_introduction.txt")
    add_section(doc, "2. Related Work", "Heading 1", "Normal", SEC / "03_related_work.txt")

    add_section(doc, "3. Method", "Heading 1", "Normal", SEC / "04_0_method_intro.txt")
    add_figure(doc, ROOT / "methodology_flow.png",
               "Figure 1. Overall methodology flow: from terrain generation "
               "(DEM + procedural fine detail), to the choice of blending function, "
               "to analytic/CPU evaluation, and to the real Isaac Lab/SAC "
               "training-analysis pipeline.",
               width_inches=6.3)
    add_section(doc, "3.1 Experimental Environment and Task Definition", "Heading 2", "Normal", SEC / "04_1_environment.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "04_1b_sim_screenshots.txt")
    add_figure(doc, ROOT / "sim_scene_view.png",
               "Figure 2. Scene view captured from running the trained policy in the "
               "real Isaac Sim environment: the four-legged lander on cratered/rocky, "
               "grainy-textured lunar terrain (high-detail local patch + directional "
               "light for documentation; the vehicle is fully vertical as verified by "
               "simulation data, up_z = +1.000).",
               width_inches=5.5)
    add_figure(doc, ROOT / "sim_lander_closeup.png",
               "Figure 3. Close-up of the same scene on the lander: the body, the tips "
               "of the gimbal thrust axis, and the centimetre-scale grainy texture of "
               "the ground are visible together.",
               width_inches=3.2)
    add_section(doc, "3.2 Simulation Setup", "Heading 2", "Normal", SEC / "04_1c_simulation_setup.txt")
    add_table(doc, TBL / "table0_sim_config.csv",
              "Table 1. Simulation, vehicle, RL training, and terrain-blending "
              "parameters shared by the analytic protocol and the real Isaac Lab/"
              "Isaac Sim training runs. Only the last group (blending parameters) "
              "varies from method to method.")
    add_section(doc, "3.3 Terrain Representation and Blending Functions", "Heading 2", "Normal", SEC / "04_2_terrain_blending.txt")
    add_section(doc, "3.4 Analytic (CPU) Evaluation Protocol", "Heading 2", "Normal", SEC / "04_3_analytic_protocol.txt")
    add_section(doc, "3.5 Reinforcement Learning Method (MDP Formulation and SAC)", "Heading 2", "Normal", SEC / "04_3b_rl_method.txt")
    add_section(doc, "3.6 Isaac Lab/Isaac Sim Training Protocol", "Heading 2", "Normal", SEC / "04_4_rl_protocol.txt")

    add_section(doc, "4. Results", "Heading 1", "Normal", SEC / "05_0_results_intro.txt")

    add_section(doc, "4.1 Boundary Continuity and Computational Cost", "Heading 2", "Normal", SEC / "05_1a_table1_para.txt")
    add_table(doc, TBL / "table1_continuity.csv",
              "Table 2. Boundary continuity ratios (relative to smoothstep), "
              "aggregated over thirty seeds, isolating only the fine-detail component.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1b_figure_radial_para.txt")
    add_figure(doc, RES / "radial_profile.png", "Figure 4. Radial fade profile of the three blending methods.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1c_table2_para.txt")
    add_table(doc, TBL / "table2_fadecost.csv",
              "Table 3. Per-query computational cost of the fade formula.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1d_figure_continuity_para.txt")
    add_figure(doc, RES / "continuity_comparison.png", "Figure 5. Cross-method comparison of C0/C1 boundary continuity.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1e_table3_para.txt")
    add_table(doc, TBL / "table3_vram.csv",
              "Table 4. Estimated VRAM usage by mesh accounting scheme.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1f_figure_cost_para.txt")
    add_figure(doc, RES / "cost_comparison.png", "Figure 6. Cross-method comparison of computational cost.")

    add_section(doc, "4.2 Ablation Analyses", "Heading 2", "Normal", SEC / "05_2a_intro.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2b_rc_sweep.txt")
    add_figure(doc, RES / "ablations" / "rc_sweep.png", "Figure 7. Contact-radius (rc) sweep: switch internal continuity and cost.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2c_switch_width.txt")
    add_figure(doc, RES / "ablations" / "switch_width_sweep.png", "Figure 8. Switch-width (w) sweep: internal continuity.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2d_radius_sweep.txt")
    add_figure(doc, RES / "ablations" / "radius_sweep.png", "Figure 9. Transition-radius (R) sweep: scale invariance.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2e_resolution_sweep.txt")
    add_figure(doc, RES / "ablations" / "resolution_sweep.png", "Figure 10. Grid-resolution sweep: empirical convergence order.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2f_wrapup.txt")

    add_section(doc, "4.3 Isaac Sim Training Results", "Heading 2", "Normal", SEC / "05_3a_walltime.txt")
    add_table(doc, TBL / "table4_rl_summary.csv",
              "Table 5. Final ten-percent window summary in the real Isaac Sim "
              "training (three seeds per method).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3b_success_summary.txt")
    add_table(doc, TBL / "table5_rl_stats.csv",
              "Table 6. Pairwise comparisons over the final ten-percent window "
              "(Mann-Whitney U, descriptive; see the caveat in Section 3.6).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3c_figures_pointer.txt")
    add_figure(doc, RES_TC / "training_comparison_metrics.png",
               "Figure 11. Outcome metrics tracked over training for the three "
               "methods (success/timeout/harsh-landing rate).")
    add_figure(doc, RES_TC / "training_comparison_diagnostics.png",
               "Figure 12. Complementary diagnostics tracked over training for the "
               "three methods (curriculum difficulty, entropy coefficient, "
               "xy-progress reward).")
    add_figure(doc, RES_TC / "final_window_boxplot.png", "Figure 13. Final ten-percent window distributions (box plot).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3d_n1_vs_n3.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3e_sample_efficiency.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3f_curriculum.txt")

    add_section(doc, "5. Discussion", "Heading 1", "Normal", SEC / "06_discussion.txt")
    add_section(doc, "6. Conclusion", "Heading 1", "Normal", SEC / "07_conclusion.txt")

    add_section(doc, "Acknowledgements", "Heading 1", "Normal", SEC / "09_acknowledgements.txt")
    add_section(doc, "Author Contributions (CRediT)", "Heading 1", "Normal", SEC / "10_author_contributions.txt")
    add_section(doc, "Conflict of Interest Statement", "Heading 1", "Normal", SEC / "11_conflict_of_interest.txt")
    add_section(doc, "Data and Code Availability", "Heading 1", "Normal", SEC / "12_data_availability.txt")

    add_references(doc, SEC / "08_references.txt")

    cp = doc.core_properties
    cp.title = title
    cp.author = "Haktan Yusuf Guloglu"
    cp.subject = ("Comparison of terrain blend methods in an Isaac Lab-based "
                  "lunar landing simulation")
    cp.language = "en-US"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
