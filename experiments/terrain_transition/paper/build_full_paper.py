#!/usr/bin/env python3
"""Builds the full paper docx in ONE pass with python-docx directly, so tables
and figures land inline in their correct subsection instead of clustering at
the end -- docx_builder.py's `merge` command drops image relationships from
any part that already contains a figure/table (verified empirically), so
multi-stage merge-then-add-figure does not work for interleaved placement.
This script reimplements the same paragraph-styling logic as
docx_builder.py's make-title-block / make-section / add-table / add-figure,
using the same style names, but on one continuously open Document object.
"""
import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

ROOT = Path("/home/haktan/Projects/LunarRocket/experiments/terrain_transition/paper")
SEC = ROOT / "sections"
TBL = ROOT / "tables"
RES = ROOT.parent / "results"
RES_TC = RES / "training_comparison"
TPL = ROOT / "generic_template.docx"
OUT = ROOT / "out" / "full_paper.docx"


def split_body_text(raw):
    for block in raw.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("- "):
            yield "bullet", block[2:].strip()
        else:
            yield "para", " ".join(line.strip() for line in block.splitlines())


def add_title_block(doc, title, authors, affiliations, email, abstract_path, keywords):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(authors)
    r.font.size = Pt(12)

    for aff in affiliations:
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(aff)
        r.italic = True
        r.font.size = Pt(10)

    doc.add_paragraph(style="Normal")

    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"*({email}) Email of the corresponding author")
    r.italic = True
    r.font.size = Pt(10)

    doc.add_paragraph(style="Normal")
    doc.add_paragraph(style="Normal")

    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract –")
    r.bold = True
    r.italic = True
    r.font.size = Pt(12)
    p.add_run(" ")
    abstract_text = Path(abstract_path).read_text(encoding="utf-8").strip()
    r = p.add_run(abstract_text)
    r.font.size = Pt(12)
    n_words = len(abstract_text.split())
    print(f"abstract word count: {n_words}")

    doc.add_paragraph(style="Normal")

    p = doc.add_paragraph(style="Normal")
    r = p.add_run("Keywords –")
    r.italic = True
    r.font.size = Pt(10)
    p.add_run(" ")
    r = p.add_run(", ".join(keywords))
    r.italic = True
    r.font.size = Pt(10)


def add_section(doc, heading, heading_style, body_style, path):
    if heading:
        doc.add_paragraph(heading, style=heading_style)
    raw = Path(path).read_text(encoding="utf-8")
    for kind, text in split_body_text(raw):
        prefix = "•  " if kind == "bullet" else ""
        doc.add_paragraph(prefix + text, style=body_style)


def add_table(doc, csv_path, caption, caption_style="Caption", table_style="Table Grid"):
    doc.add_paragraph(caption, style=caption_style)
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = table_style
    except KeyError:
        pass
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val


def add_figure(doc, image_path, caption, caption_style="Caption", width_inches=6.2):
    doc.add_picture(str(image_path), width=Inches(width_inches) if width_inches else None)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(caption, style=caption_style)


def add_references(doc, path):
    from docx.shared import Cm
    doc.add_paragraph("Kaynaklar", style="Heading 1")
    raw = Path(path).read_text(encoding="utf-8")
    for block in raw.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        text = " ".join(line.strip() for line in block.splitlines())
        p = doc.add_paragraph(text, style="Normal")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)


def main():
    doc = Document(str(TPL))
    body = doc.element.body
    from docx.oxml.ns import qn
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    add_title_block(
        doc,
        title="Ay İnişi Simülasyonlarında Arazi İnce-Detay Geçişi: Isaac Lab'da Smoothstep, Gauss ve Hibrit Karıştırma Yöntemlerinin Karşılaştırılması",
        authors="Haktan Yusuf Güloğlu",
        affiliations=["Sakarya Üniversitesi, Millî Teknoloji Atölyesi, Yapay Zeka Merkezi"],
        email="haktanguloglu24@gmail.com",
        abstract_path=ROOT / "abstract_en.txt",
        keywords=["Lunar landing simulation", "Reinforcement learning", "Isaac Lab",
                  "Procedural terrain generation", "Multi-resolution blending",
                  "Continuity analysis", "Soft Actor-Critic"],
    )

    add_section(doc, "Özet", "Heading 1", "Normal", SEC / "01_ozet_tr.txt")
    add_section(doc, "1. Giriş", "Heading 1", "Normal", SEC / "02_giris.txt")
    add_section(doc, "2. İlgili Çalışmalar", "Heading 1", "Normal", SEC / "03_ilgili_calismalar.txt")

    add_section(doc, "3. Yöntem", "Heading 1", "Normal", SEC / "04_0_yontem_giris.txt")
    add_figure(doc, ROOT / "methodology_flow.png",
               "Şekil 1. Genel metodoloji akış şeması: arazi üretiminden (DEM + prosedürel ince detay), karıştırma fonksiyonu seçimine, analitik/CPU değerlendirmesine ve gerçek Isaac Lab/SAC eğitim-analiz hattına.",
               width_inches=6.3)
    add_section(doc, "3.1 Deney Ortamı ve Görev Tanımı", "Heading 2", "Normal", SEC / "04_1_deney_ortami.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "04_1b_sim_screenshots.txt")
    add_figure(doc, ROOT / "sim_scene_view.png",
               "Şekil 2. Eğitilmiş politikanın gerçek Isaac Sim ortamında çalıştırılmasından alınmış sahne görüntüsü: dört bacaklı iniş aracı, kraterli/kayalık ve taneli dokulu ay zemini üzerinde (yüksek-detay yerel yama + belgeleme amaçlı yönlü ışık; araç simülasyon verisiyle doğrulanmış biçimde tam dikeydir, up_z=+1,000).",
               width_inches=5.5)
    add_figure(doc, ROOT / "sim_lander_closeup.png",
               "Şekil 3. Aynı sahnenin iniş aracına yakınlaştırılmış görünümü: gövde, gimbal itki ekseninin uçları ve zeminin santimetre ölçekli taneli dokusu birlikte görülebilmektedir.",
               width_inches=3.2)
    add_section(doc, "3.2 Simülasyon Kurulumu", "Heading 2", "Normal", SEC / "04_1c_simulasyon_kurulumu.txt")
    add_table(doc, TBL / "table0_sim_config.csv",
              "Tablo 1. Analitik protokol ile gerçek Isaac Lab/Isaac Sim eğitim koşularının paylaştığı simülasyon, araç, RL eğitim ve arazi karıştırma parametreleri. Yalnızca son grup (karıştırma parametreleri), yöntemden yönteme değişir.")
    add_section(doc, "3.3 Arazi Temsili ve Karıştırma Fonksiyonları", "Heading 2", "Normal", SEC / "04_2_arazi_karistirma.txt")
    add_section(doc, "3.4 Analitik (CPU) Değerlendirme Protokolü", "Heading 2", "Normal", SEC / "04_3_analitik_protokol.txt")
    add_section(doc, "3.5 Deneyde Kullanılan Pekiştirmeli Öğrenme Yöntemi (MDP Formülasyonu ve SAC)", "Heading 2", "Normal", SEC / "04_3b_rl_yontem.txt")
    add_section(doc, "3.6 Gerçek Isaac Lab/Isaac Sim Eğitim Protokolü", "Heading 2", "Normal", SEC / "04_4_rl_protokol.txt")

    add_section(doc, "4. Bulgular", "Heading 1", "Normal", SEC / "05_0_bulgular_giris.txt")

    add_section(doc, "4.1 Sınır Sürekliliği ve Hesaplama Maliyeti", "Heading 2", "Normal", SEC / "05_1a_tablo1_para.txt")
    add_table(doc, TBL / "table1_continuity.csv",
              "Tablo 2. Otuz tohum üzerinde toplanmış, yalnızca ince-detay bileşenini izole eden sınır sürekliliği oranları (smoothstep'e göre).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1b_figure_radial_para.txt")
    add_figure(doc, RES / "radial_profile.png", "Şekil 4. Üç karıştırma yönteminin radyal söndürme profili.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1c_tablo2_para.txt")
    add_table(doc, TBL / "table2_fadecost.csv",
              "Tablo 3. Söndürme (fade) formülünün sorgu başına hesaplama maliyeti.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1d_figure_continuity_para.txt")
    add_figure(doc, RES / "continuity_comparison.png", "Şekil 5. Yöntemler arası C0/C1 sınır sürekliliği karşılaştırması.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1e_tablo3_para.txt")
    add_table(doc, TBL / "table3_vram.csv",
              "Tablo 4. Ağ muhasebe biçimine göre tahmini VRAM kullanımı.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_1f_figure_cost_para.txt")
    add_figure(doc, RES / "cost_comparison.png", "Şekil 6. Yöntemler arası hesaplama maliyeti karşılaştırması.")

    add_section(doc, "4.2 Ablasyon Analizleri", "Heading 2", "Normal", SEC / "05_2a_intro.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2b_rc_sweep.txt")
    add_figure(doc, RES / "ablations" / "rc_sweep.png", "Şekil 7. Temas yarıçapı (rc) taraması: switch iç sürekliliği ve maliyet.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2c_switch_width.txt")
    add_figure(doc, RES / "ablations" / "switch_width_sweep.png", "Şekil 8. Switch genişliği (w) taraması: iç süreklilik.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2d_radius_sweep.txt")
    add_figure(doc, RES / "ablations" / "radius_sweep.png", "Şekil 9. Geçiş yarıçapı (R) taraması: ölçek değişmezliği.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2e_resolution_sweep.txt")
    add_figure(doc, RES / "ablations" / "resolution_sweep.png", "Şekil 10. Grid çözünürlüğü taraması: ampirik yakınsama derecesi.")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_2f_wrapup.txt")

    add_section(doc, "4.3 Gerçek Isaac Sim Eğitim Sonuçları", "Heading 2", "Normal", SEC / "05_3a_walltime.txt")
    add_table(doc, TBL / "table4_rl_summary.csv",
              "Tablo 5. Gerçek Isaac Sim eğitiminde son yüzde onluk pencere özeti (yöntem başına üç tohum).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3b_success_summary.txt")
    add_table(doc, TBL / "table5_rl_stats.csv",
              "Tablo 6. Son yüzde onluk pencere üzerinde ikili karşılaştırmalar (Mann-Whitney U, betimsel; bkz. Bölüm 3.6 çekincesi).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3c_figures_pointer.txt")
    add_figure(doc, RES_TC / "training_comparison_metrics.png",
               "Şekil 11. Üç yöntem için eğitim boyunca izlenen sonuç metrikleri (başarı/zaman aşımı/sert iniş oranı).")
    add_figure(doc, RES_TC / "training_comparison_diagnostics.png",
               "Şekil 12. Üç yöntem için eğitim boyunca izlenen tamamlayıcı diyagnostikler (müfredat zorluğu, entropi katsayısı, xy-ilerleme ödülü).")
    add_figure(doc, RES_TC / "final_window_boxplot.png", "Şekil 13. Son yüzde onluk pencere dağılımları (kutu grafiği).")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3d_n1_vs_n3.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3e_sample_efficiency.txt")
    add_section(doc, None, "Heading 2", "Normal", SEC / "05_3f_curriculum.txt")

    add_section(doc, "5. Tartışma", "Heading 1", "Normal", SEC / "06_tartisma.txt")
    add_section(doc, "6. Sonuç", "Heading 1", "Normal", SEC / "07_sonuc.txt")

    add_section(doc, "Teşekkür", "Heading 1", "Normal", SEC / "09_tesekkur.txt")
    add_section(doc, "Yazar Katkı Beyanı (CRediT)", "Heading 1", "Normal", SEC / "10_yazar_katki.txt")
    add_section(doc, "Çıkar Çatışması Beyanı", "Heading 1", "Normal", SEC / "11_cikar_catismasi.txt")
    add_section(doc, "Veri ve Kod Erişilebilirliği", "Heading 1", "Normal", SEC / "12_veri_erisebilirlik.txt")

    add_references(doc, SEC / "08_kaynaklar.txt")

    cp = doc.core_properties
    cp.title = "Ay İnişi Simülasyonlarında Arazi İnce-Detay Geçişi: Isaac Lab'da Smoothstep, Gauss ve Hibrit Karıştırma Yöntemlerinin Karşılaştırılması"
    cp.author = "Haktan Yusuf Güloğlu"
    cp.subject = "Isaac Lab tabanlı ay inişi simülasyonunda terrain blend yöntemlerinin karşılaştırılması"
    cp.language = "tr-TR"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
